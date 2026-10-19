# -*- coding: utf-8 -*-
"""Extra deliverables: editable manuscript (.md + .docx), EPUB eBook,
and a compressed screen/web PDF."""
import io
import re
import uuid
import zipfile
from pathlib import Path
from html import escape as hesc

from PIL import Image as PILImage

from book_content import (BOOK_TITLE, BOOK_SUBTITLE, ILLUSTRATOR, DEDICATION,
                          DEDICATION_TO, INTRO_TITLE, INTRO_PAGES, STORIES)

ROOT = Path(__file__).parent
AST = ROOT / "book_assets"
OUT = ROOT / "output"
MAIN_PDF = OUT / "Old_Old_Tales_Retold_Restored_Edition.pdf"


def clean(text):
    return re.sub(r"\[\[(?:BIG|MED|WEE):(.*?)\]\]", r"\1", text)


def bear_html(text):
    def repl(m):
        cls = {"BIG": "bear-big", "MED": "bear-med", "WEE": "bear-wee"}[m.group(1)]
        return f'<span class="{cls}">{hesc(m.group(2))}</span>'
    return re.sub(r"\[\[(BIG|MED|WEE):(.*?)\]\]", repl, hesc(text.replace("&", "&amp;")) if False else text)


# ---------------- manuscript ----------------
def manuscript_md():
    lines = [f"# {BOOK_TITLE}", "", f"*{BOOK_SUBTITLE}*", "",
             f"Illustrated by {ILLUSTRATOR}", "", "---", "",
             f"*{DEDICATION} {DEDICATION_TO}.*", "", "---", "",
             f"## {INTRO_TITLE}", ""]
    for paras in INTRO_PAGES:
        for p in paras:
            lines += [clean(p), ""]
    for s in STORIES:
        lines += [f"## {s['title']}", ""]
        for blocks, _plate in s["pages"]:
            for b in blocks:
                if b.startswith("> "):
                    for ln in b.split("\n"):
                        lines.append("> " + clean(ln[2:]))
                    lines.append("")
                else:
                    lines += [clean(b), ""]
    lines += ["---", "", "*The End*", ""]
    path = OUT / "Manuscript_Old_Old_Tales_Retold.md"
    path.write_text("\n".join(lines), encoding="utf-8")
    return path


def manuscript_docx():
    from docx import Document
    from docx.shared import Pt, Inches
    from docx.enum.text import WD_ALIGN_PARAGRAPH

    doc = Document()
    for sec in doc.sections:
        sec.top_margin = sec.bottom_margin = Inches(1)
        sec.left_margin = sec.right_margin = Inches(1.1)
    style = doc.styles["Normal"]
    style.font.name = "Georgia"
    style.font.size = Pt(12)
    style.paragraph_format.space_after = Pt(8)
    style.paragraph_format.line_spacing = 1.35

    t = doc.add_heading(BOOK_TITLE, level=0)
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph(BOOK_SUBTITLE); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].italic = True
    p = doc.add_paragraph(f"Illustrated by {ILLUSTRATOR}"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph("")
    p = doc.add_paragraph(f"{DEDICATION} {DEDICATION_TO}.")
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.runs[0].italic = True
    doc.add_page_break()

    doc.add_heading(INTRO_TITLE, level=1)
    for paras in INTRO_PAGES:
        for para in paras:
            doc.add_paragraph(clean(para))

    for s in STORIES:
        doc.add_page_break()
        doc.add_heading(s["title"], level=1)
        for blocks, _plate in s["pages"]:
            for b in blocks:
                if b.startswith("> "):
                    for ln in b.split("\n"):
                        p = doc.add_paragraph(clean(ln[2:]))
                        p.paragraph_format.left_indent = Inches(0.45)
                        p.paragraph_format.space_after = Pt(2)
                    doc.add_paragraph("").paragraph_format.space_after = Pt(2)
                else:
                    doc.add_paragraph(clean(b))
    doc.add_page_break()
    p = doc.add_paragraph("The End"); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.runs[0].bold = True

    doc.add_heading("Editor's Notes (not part of the book)", level=1)
    for note in [
        "Transcribed page-by-page from photographs of the 1923 Volland Edition; wording is faithful to that copy.",
        "Obvious printer's errors in the old copy were quietly corrected (a dropped 'to' in Chicken Licken, 'So the went' → 'So they went', one missing closing quote in Little Tuppens, and a few missing sentence stops).",
        "This copy's text differs in a few places from the later Donohue reprint: in this original, the wolf eats the first two pigs and is burned in the chimney, Little Tuppens is 'choking' (not coughing), and the cat 'kills' the rat / the rope 'hangs' the butcher. Soften these if you prefer before publishing.",
        "ISBN and publisher details are placeholders on the copyright page and back cover of the PDF — tell your formatter what to insert.",
    ]:
        doc.add_paragraph(note, style="List Bullet")
    path = OUT / "Manuscript_Old_Old_Tales_Retold.docx"
    doc.save(path)
    return path


# ---------------- EPUB ----------------
CSS = """
body{font-family:Georgia,'Times New Roman',serif;line-height:1.55;color:#2B2118;margin:5%;}
h1,h2{font-weight:normal;text-align:center;line-height:1.25;}
h1{font-size:1.9em;margin:0.4em 0 0.2em;}
h2{font-size:1.5em;margin:0.3em 0 0.6em;}
p{text-indent:1.2em;margin:0 0 0.55em;text-align:left;}
p.first,p.dlg{text-indent:0;}
p.center{text-align:center;text-indent:0;}
p.verse{text-indent:0;margin-left:2.2em;margin-bottom:0.7em;}
figure{margin:0.8em 0;text-align:center;}
figure img{max-width:100%;}
.title-plate img{max-width:96%;}
.kicker{text-align:center;letter-spacing:0.18em;font-size:0.78em;color:#8A5A2B;text-indent:0;}
.orn{text-align:center;color:#8A5A2B;text-indent:0;letter-spacing:0.4em;}
.bear-big{font-size:1.18em;} .bear-med{font-size:1.02em;} .bear-wee{font-size:0.82em;}
.small{font-size:0.85em;color:#5A4A38;}
"""


def xhtml(title, body):
    return ('<?xml version="1.0" encoding="utf-8"?>\n'
            '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
            f'<head><title>{hesc(title)}</title><link rel="stylesheet" href="style.css"/></head>'
            f'<body>{body}</body></html>')


def para_html(block, first=False):
    if block.startswith("> "):
        lines = [bear_html(ln[2:]) for ln in block.split("\n")]
        return '<p class="verse">' + "<br/>".join(lines) + "</p>"
    cls = "first" if first else ("dlg" if block.lstrip().startswith("“") else "")
    return f'<p class="{cls}">{bear_html(block)}</p>'


def img_bytes(path, max_w=900, quality=80):
    im = PILImage.open(path).convert("RGB")
    if im.width > max_w:
        im = im.resize((max_w, int(im.height * max_w / im.width)), PILImage.LANCZOS)
    buf = io.BytesIO(); im.save(buf, "JPEG", quality=quality)
    return buf.getvalue()


def build_epub(cover_jpg: Path):
    uid = f"urn:uuid:{uuid.uuid4()}"
    files = {}          # arcname -> bytes/str
    manifest, spine, navs = [], [], []

    def add(fid, href, media, content, props=None, in_spine=True):
        files[href] = content
        manifest.append((fid, href, media, props))
        if in_spine:
            spine.append(fid)

    files["style.css"] = CSS
    manifest.append(("css", "style.css", "text/css", None))

    # images
    plates = sorted({s["title_plate"] for s in STORIES} | {p for s in STORIES for _, p in s["pages"]})
    for pn in plates:
        href = f"images/plate-{pn:03d}.jpg"
        files[href] = img_bytes(AST / "plates" / f"plate-{pn:03d}.jpg")
        manifest.append((f"img{pn}", href, "image/jpeg", None))
    files["images/cover.jpg"] = img_bytes(cover_jpg, max_w=1000, quality=85)
    manifest.append(("cover-img", "images/cover.jpg", "image/jpeg", "cover-image"))

    # cover + title
    add("cover", "cover.xhtml", "application/xhtml+xml",
        xhtml(BOOK_TITLE, '<figure><img src="images/cover.jpg" alt="Cover"/></figure>'))
    add("title", "title.xhtml", "application/xhtml+xml", xhtml(BOOK_TITLE,
        f'<h1>{hesc(BOOK_TITLE)}</h1><p class="center"><i>{hesc(BOOK_SUBTITLE)}</i></p>'
        f'<p class="orn">&#9670;</p><p class="center">Illustrated by<br/>{hesc(ILLUSTRATOR)}</p>'
        f'<p class="center small">THE VOLLAND EDITION &#8226; RESTORED EDITION</p>'
        f'<p class="center"><i>{hesc(DEDICATION)} {hesc(DEDICATION_TO)}.</i></p>'))
    navs.append((INTRO_TITLE, "intro.xhtml"))

    intro_body = f"<h2>{hesc(INTRO_TITLE)}</h2>"
    first = True
    for paras in INTRO_PAGES:
        for p in paras:
            intro_body += para_html(p, first=first); first = False
    add("intro", "intro.xhtml", "application/xhtml+xml", xhtml(INTRO_TITLE, intro_body))

    words = ["One", "Two", "Three", "Four", "Five", "Six", "Seven", "Eight"]
    for idx, s in enumerate(STORIES):
        fid = f"story{idx+1}"; href = f"story{idx+1}.xhtml"
        navs.append((s["title"], href))
        body = (f'<p class="kicker">STORY {words[idx].upper()}</p><h2>{hesc(s["title"])}</h2>'
                f'<figure class="title-plate"><img src="images/plate-{s["title_plate"]:03d}.jpg" alt=""/></figure>')
        for blocks, plate_no in s["pages"]:
            for i, b in enumerate(blocks):
                body += para_html(b, first=(i == 0))
            body += f'<figure><img src="images/plate-{plate_no:03d}.jpg" alt=""/></figure>'
        body += '<p class="orn">&#9670;</p>'
        add(fid, href, "application/xhtml+xml", xhtml(s["title"], body))

    add("end", "end.xhtml", "application/xhtml+xml", xhtml("The End",
        '<h2>The End</h2><p class="center"><i>And so end the old, old tales — told over and over, '
        'and loved by children everywhere.</i></p>'))

    nav_items = "".join(f'<li><a href="{h}">{hesc(t)}</a></li>' for t, h in navs)
    nav_doc = ('<?xml version="1.0" encoding="utf-8"?>\n'
               '<html xmlns="http://www.w3.org/1999/xhtml" xmlns:epub="http://www.idpf.org/2007/ops">'
               f'<head><title>Contents</title><link rel="stylesheet" href="style.css"/></head>'
               f'<body><nav epub:type="toc"><h2>Contents</h2><ol>{nav_items}</ol></nav></body></html>')
    files["nav.xhtml"] = nav_doc
    manifest.append(("nav", "nav.xhtml", "application/xhtml+xml", "nav"))

    ncx_items = "".join(
        f'<navPoint id="n{i}" playOrder="{i}"><navLabel><text>{hesc(t)}</text></navLabel>'
        f'<content src="{h}"/></navPoint>' for i, (t, h) in enumerate(navs, 1))
    files["toc.ncx"] = ('<?xml version="1.0" encoding="utf-8"?>\n'
        f'<ncx xmlns="http://www.daisy.org/z3986/2005/ncx/" version="2005-1"><head>'
        f'<meta name="dtb:uid" content="{uid}"/></head>'
        f'<docTitle><text>{hesc(BOOK_TITLE)}</text></docTitle><navMap>{ncx_items}</navMap></ncx>')
    manifest.append(("ncx", "toc.ncx", "application/x-dtbncx+xml", None))

    man = "".join(f'<item id="{i}" href="{h}" media-type="{m}"' + (f' properties="{pr}"' if pr else "") + "/>"
                  for i, h, m, pr in manifest)
    spn = "".join(f'<itemref idref="{i}"/>' for i in spine)
    opf = ('<?xml version="1.0" encoding="utf-8"?>\n'
           '<package xmlns="http://www.idpf.org/2007/opf" unique-identifier="bk" version="3.0">'
           '<metadata xmlns:dc="http://purl.org/dc/elements/1.1/">'
           f'<dc:identifier id="bk">{uid}</dc:identifier>'
           f'<dc:title>{hesc(BOOK_TITLE)}: {hesc(BOOK_SUBTITLE)}</dc:title>'
           '<dc:language>en</dc:language>'
           f'<dc:creator>{hesc(ILLUSTRATOR)} (illustrator)</dc:creator>'
           '<dc:description>Eight best-beloved folk stories for children, restored from the 1923 Volland Edition.</dc:description>'
           '</metadata>'
           f'<manifest>{man}</manifest><spine toc="ncx">{spn}</spine></package>')
    container = ('<?xml version="1.0" encoding="utf-8"?>\n'
                 '<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">'
                 '<rootfiles><rootfile full-path="content.opf" media-type="application/oebps-package+xml"/></rootfiles></container>')

    out = OUT / "Old_Old_Tales_Retold_Restored_Edition.epub"
    with zipfile.ZipFile(out, "w") as z:
        z.writestr(zipfile.ZipInfo("mimetype"), "application/epub+zip", zipfile.ZIP_STORED)
        z.writestr("META-INF/container.xml", container)
        z.writestr("content.opf", opf)
        for name, content in files.items():
            z.writestr(name, content, zipfile.ZIP_DEFLATED)
    return out


def build_web_pdf():
    import fitz
    doc = fitz.open(MAIN_PDF)
    try:
        doc.rewrite_images(dpi_threshold=150, dpi_target=115, quality=72)
    except Exception as e:
        print("rewrite_images skipped:", e)
    out = OUT / "Old_Old_Tales_Retold_Web.pdf"
    doc.save(out, garbage=4, deflate=True, clean=True)
    doc.close()
    return out


def render_cover():
    import fitz
    doc = fitz.open(MAIN_PDF)
    p = ROOT / "book_assets" / "cover_render.jpg"
    doc[0].get_pixmap(dpi=170).save(p)
    doc.close()
    return p


if __name__ == "__main__":
    print(manuscript_md())
    print(manuscript_docx())
    cover = render_cover()
    print(build_epub(cover))
    print(build_web_pdf())
